from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = r'D:\VScode\JavaWeb\web-project02\Java面试八股与手撕代码复习资料.docx'
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.75); sec.right_margin = Inches(.75)

styles = doc.styles
styles['Normal'].font.name = 'Microsoft YaHei'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei'); styles['Normal'].font.size = Pt(10.5)
for name, size, color in [('Title', 24, '000000'), ('Heading 1', 16, '1F4E79'), ('Heading 2', 12, '2F5597')]:
    st = styles[name]; st.font.name='Microsoft YaHei'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(color)
styles['Title'].paragraph_format.space_after = Pt(12)
styles['Heading 1'].paragraph_format.space_before = Pt(12); styles['Heading 1'].paragraph_format.space_after = Pt(5)
styles['Heading 2'].paragraph_format.space_before = Pt(8); styles['Heading 2'].paragraph_format.space_after = Pt(3)

code = styles.add_style('Code Block', WD_STYLE_TYPE.PARAGRAPH)
code.font.name='Consolas'; code._element.rPr.rFonts.set(qn('w:eastAsia'),'Consolas'); code.font.size=Pt(9); code.font.color.rgb=RGBColor.from_string('333333')
code.paragraph_format.left_indent=Inches(.25); code.paragraph_format.space_after=Pt(5)

def p(text='', style=None, bold_prefix=None):
    para=doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        para.add_run(bold_prefix).bold=True; para.add_run(text[len(bold_prefix):])
    else: para.add_run(text)
    return para
def bullets(items):
    for x in items: p(x, 'List Bullet')
def q(title, answer):
    p(title, 'Heading 2'); p(answer)
def codeblock(text):
    for line in text.strip('\n').split('\n'): p(line, 'Code Block')

title=doc.add_paragraph(style='Title'); title.alignment=WD_ALIGN_PARAGRAPH.CENTER; title.add_run('Java 面试八股与手撕代码复习资料')
sub=doc.add_paragraph(); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER; sub.add_run('高频知识点 标准回答与常见代码题').italic=True
p('本资料面向 Java 后端实习和校招面试复习，覆盖 Java 基础、集合、JVM、并发、MySQL、Spring、Redis，以及面试高频手撕代码。回答建议采用“先结论、再原理、最后结合场景”的结构。')

doc.add_heading('一 Java 基础', level=1)
q('1 == 和 equals 有什么区别', '== 比较基本类型的值，比较引用类型时比较对象地址。equals 默认也是比较地址，但 String 等类重写了 equals，用于比较内容。重写 equals 时通常必须同时重写 hashCode，否则放入 HashMap 或 HashSet 可能出现问题。')
q('2 重载和重写的区别', '重载发生在同一个类中，方法名相同但参数列表不同，是编译期多态。重写发生在父子类之间，子类重新实现父类方法，是运行期多态。重写时访问权限不能更严格。')
q('3 String 为什么不可变', '不可变有利于字符串常量池复用、线程安全、hashCode 缓存和安全性，也使 String 适合作为 HashMap 的 key。')
q('4 String StringBuilder StringBuffer 的区别', 'String 不可变；StringBuilder 可变但线程不安全，性能较高；StringBuffer 可变且方法带同步，线程安全但性能相对较低。单线程拼接通常使用 StringBuilder。')
q('5 final finally finalize 的区别', 'final 用于修饰变量、方法和类；finally 是异常处理中的代码块，通常用于释放资源；finalize 是对象回收前可能调用的方法，已经不推荐使用。')
q('6 Java 是值传递还是引用传递', 'Java 只有值传递。传递对象时，传递的是引用变量的副本，因此可以通过副本修改对象内容，但不能让原引用指向另一个对象。')

doc.add_heading('二 集合', level=1)
q('7 ArrayList 和 LinkedList 的区别', 'ArrayList 底层是动态数组，随机访问快，中间插入和删除需要移动元素。LinkedList 是双向链表，随机访问慢，但节点插入删除不需要整体移动。实际业务中大多数场景优先使用 ArrayList。')
q('8 HashMap 底层原理', 'JDK 8 中 HashMap 底层是数组、链表和红黑树。put 时先扰动 hash，再计算数组下标；位置为空则直接插入，发生冲突时使用链表或红黑树。元素超过负载因子阈值后扩容，默认负载因子为 0.75。')
q('9 HashMap 为什么线程不安全', '并发修改可能造成数据覆盖、size 不准确、扩容数据丢失或读到不一致数据。并发场景应使用 ConcurrentHashMap。')
q('10 HashSet 如何保证不重复', 'HashSet 底层基于 HashMap，元素作为 key 保存。添加时先比较 hashCode，再比较 equals，两者都相同则认为重复。')

doc.add_heading('三 JVM', level=1)
q('11 JVM 内存区域有哪些', '程序计数器记录线程执行位置；虚拟机栈保存局部变量和方法调用信息；本地方法栈执行 Native 方法；堆保存对象实例，是垃圾回收的主要区域；方法区或元空间保存类信息、常量和静态变量。')
q('12 如何判断对象是否应该回收', '主要使用可达性分析：从 GC Roots 出发，如果对象无法通过引用链到达，就认为不可达，可以被回收。GC Roots 包括栈中的局部变量、静态变量、常量引用和活动线程对象等。')
q('13 什么是内存泄漏', '对象已经没有使用价值，但仍被其他对象引用，导致 GC 无法回收。常见原因包括静态集合、ThreadLocal 使用不当、未关闭资源、监听器未移除和无过期缓存。')

doc.add_heading('四 并发编程', level=1)
q('14 sleep wait join 的区别', 'sleep 是 Thread 的方法，不释放锁；wait 是 Object 的方法，必须在 synchronized 中调用并释放锁，需要 notify 或 notifyAll 唤醒；join 用于等待另一个线程执行结束。')
q('15 volatile 有什么作用', 'volatile 保证可见性和一定程度的有序性，但不能保证复合操作的原子性。例如 count++ 包含读取、加一、写回三个步骤，不能仅靠 volatile 保证安全。')
q('16 synchronized 和 Lock 的区别', 'synchronized 自动加锁和释放锁，使用简单；Lock 需要手动释放，但支持可中断、超时、公平锁和多个 Condition。')
q('17 线程池执行流程', '核心线程未满时创建核心线程；核心线程满后任务进入阻塞队列；队列满后创建非核心线程；达到最大线程数后执行拒绝策略。常见策略有 AbortPolicy、CallerRunsPolicy、DiscardPolicy 和 DiscardOldestPolicy。')
q('18 ThreadLocal 为什么可能内存泄漏', 'ThreadLocalMap 的 key 是弱引用，但 value 是强引用。ThreadLocal 被回收后，value 可能在线程池线程中长期存在，使用完应在 finally 中调用 remove。')

doc.add_heading('五 MySQL 与 Spring', level=1)
q('19 MySQL 为什么使用 B+ 树', '非叶子节点只保存索引，叶子节点保存数据或数据地址且通过链表连接，因此树高度较低、磁盘 IO 少，并且适合范围查询。')
q('20 什么是回表和覆盖索引', '普通索引找到主键后还要回到聚簇索引查询完整数据，这叫回表。查询字段都包含在索引中时可以直接返回，称为覆盖索引。')
q('21 联合索引的最左匹配原则', '联合索引 (a,b,c) 可以支持按 a、a+b 或 a+b+c 查询，但不能有效支持只按 b 或 c 查询。遇到范围查询后，后面的列通常不能继续用于索引定位。')
q('22 IOC DI 和 AOP', 'IOC 是把对象创建和依赖管理交给 Spring 容器；DI 是依赖注入，是实现 IOC 的方式。AOP 通过动态代理实现，常用于事务、日志、权限和性能统计。')
q('23 @Transactional 为什么会失效', '常见原因包括方法不是 public、同类内部直接调用、异常被 catch 后没有抛出、抛出非回滚异常、Bean 没有被 Spring 管理，以及数据库引擎不支持事务。')

doc.add_heading('六 Redis', level=1)
q('24 Redis 为什么快', 'Redis 主要基于内存操作，数据结构高效，使用 IO 多路复用，避免大量线程切换和锁竞争，并采用 C 语言实现。Redis 6 以后网络 IO 可以使用多线程，但命令执行仍主要是单线程模型。')
q('25 缓存穿透 击穿 雪崩', '穿透是查询不存在的数据，可用缓存空值或布隆过滤器；击穿是热点 key 过期导致大量请求打到数据库，可用互斥锁或逻辑过期；雪崩是大量 key 同时过期或 Redis 故障，可用随机过期时间、集群、限流和降级。')

doc.add_heading('七 高频手撕代码', level=1)
p('下面代码均可直接作为面试模板。面试时先说明思路、时间复杂度和边界条件，再开始编码。')
p('1 反转链表', 'Heading 2')
codeblock('''public ListNode reverseList(ListNode head) {
    ListNode prev = null;
    ListNode curr = head;
    while (curr != null) {
        ListNode next = curr.next;
        curr.next = prev;
        prev = curr;
        curr = next;
    }
    return prev;
}''')
p('2 判断链表是否有环', 'Heading 2')
codeblock('''public boolean hasCycle(ListNode head) {
    ListNode slow = head, fast = head;
    while (fast != null && fast.next != null) {
        slow = slow.next;
        fast = fast.next.next;
        if (slow == fast) return true;
    }
    return false;
}''')
p('3 两数之和', 'Heading 2')
codeblock('''public int[] twoSum(int[] nums, int target) {
    Map<Integer, Integer> map = new HashMap<>();
    for (int i = 0; i < nums.length; i++) {
        int need = target - nums[i];
        if (map.containsKey(need)) return new int[]{map.get(need), i};
        map.put(nums[i], i);
    }
    return new int[0];
}''')
p('4 最长无重复子串', 'Heading 2')
codeblock('''public int lengthOfLongestSubstring(String s) {
    Set<Character> set = new HashSet<>();
    int left = 0, result = 0;
    for (int right = 0; right < s.length(); right++) {
        while (set.contains(s.charAt(right))) set.remove(s.charAt(left++));
        set.add(s.charAt(right));
        result = Math.max(result, right - left + 1);
    }
    return result;
}''')
p('5 二叉树层序遍历', 'Heading 2')
codeblock('''public List<List<Integer>> levelOrder(TreeNode root) {
    List<List<Integer>> result = new ArrayList<>();
    if (root == null) return result;
    Queue<TreeNode> queue = new LinkedList<>();
    queue.offer(root);
    while (!queue.isEmpty()) {
        int size = queue.size();
        List<Integer> level = new ArrayList<>();
        for (int i = 0; i < size; i++) {
            TreeNode node = queue.poll();
            level.add(node.val);
            if (node.left != null) queue.offer(node.left);
            if (node.right != null) queue.offer(node.right);
        }
        result.add(level);
    }
    return result;
}''')
p('6 有效括号', 'Heading 2')
codeblock('''public boolean isValid(String s) {
    Deque<Character> stack = new ArrayDeque<>();
    for (char c : s.toCharArray()) {
        if (c == '(' || c == '[' || c == '{') stack.push(c);
        else {
            if (stack.isEmpty()) return false;
            char t = stack.pop();
            if ((c == ')' && t != '(') || (c == ']' && t != '[')
                    || (c == '}' && t != '{')) return false;
        }
    }
    return stack.isEmpty();
}''')
p('7 二分查找', 'Heading 2')
codeblock('''public int search(int[] nums, int target) {
    int left = 0, right = nums.length - 1;
    while (left <= right) {
        int mid = left + (right - left) / 2;
        if (nums[mid] == target) return mid;
        if (nums[mid] < target) left = mid + 1;
        else right = mid - 1;
    }
    return -1;
}''')
p('8 LRU 缓存', 'Heading 2')
codeblock('''class LRUCache extends LinkedHashMap<Integer, Integer> {
    private final int capacity;
    public LRUCache(int capacity) {
        super(capacity, 0.75f, true);
        this.capacity = capacity;
    }
    public int get(int key) { return super.getOrDefault(key, -1); }
    public void put(int key, int value) { super.put(key, value); }
    protected boolean removeEldestEntry(Map.Entry<Integer, Integer> e) {
        return size() > capacity;
    }
}''')

doc.add_heading('八 消息队列与微服务', level=1)
q('26 为什么使用消息队列', '异步处理可以降低接口响应时间；应用解耦可以减少系统之间的直接依赖；削峰填谷可以缓解突发流量对下游服务和数据库的压力。')
q('27 消息丢失 重复消费和消息积压如何处理', '消息丢失可通过生产者确认、持久化、消费者手动确认和失败重试解决。重复消费要保证业务幂等，例如唯一业务号、去重表或状态机。消息积压要扩容消费者、提高消费能力、优化处理逻辑，并监控队列长度。')
q('28 什么是 CAP 和 BASE', 'CAP 指一致性、可用性和分区容错性不能同时完美满足，分布式系统通常必须保证分区容错，再在一致性和可用性之间取舍。BASE 是基本可用、软状态和最终一致性，适合高并发分布式场景。')

doc.add_heading('九 补充高频八股', level=1)
q('29 受检异常和非受检异常', '受检异常继承 Exception 但不包括 RuntimeException，编译器要求显式处理或声明抛出。非受检异常通常继承 RuntimeException，编译器不强制处理。Error 通常表示程序无法处理的严重问题。')
q('30 泛型中的 extends 和 super', ' extends 表示上界，只能安全读取； super 表示下界，适合写入。记忆原则是 PECS：Producer Extends，Consumer Super。')
q('31 双亲委派模型有什么作用', '类加载器收到加载请求后先交给父加载器，父加载器无法完成时才由子加载器处理。这样可以避免核心类被重复加载，也防止用户自定义类替换 Java 核心类。')
q('32 CAS 和 ABA 问题', 'CAS 使用期望值和内存当前值比较，只有相同时才更新，常用于无锁并发。ABA 是值从 A 变成 B 又变回 A，CAS 只看到最终仍是 A。可以通过版本号或 AtomicStampedReference 解决。')
q('33 乐观锁和悲观锁', '悲观锁认为冲突经常发生，访问数据前先加锁；乐观锁认为冲突较少，通过版本号或 CAS 更新，失败后重试。乐观锁适合读多写少，悲观锁适合冲突较多的场景。')
q('34 RDB 和 AOF 的区别', 'RDB 是定时生成数据快照，文件小、恢复快，但可能丢失最近数据。AOF 记录写命令，数据更完整，但文件通常更大。实际部署可以同时开启两者。')
q('35 Redis 分布式锁的注意事项', '加锁时使用 SET key value NX EX，value 必须唯一；释放锁时要校验 value，并用 Lua 保证判断和删除的原子性；还要考虑业务执行时间超过锁过期时间的问题。')
q('36 Spring 如何解决循环依赖', '对于单例 Bean 的属性注入循环依赖，Spring 通常通过三级缓存提前暴露早期对象解决。构造器注入的循环依赖通常无法自动解决。')
q('37 过滤器和拦截器的区别', '过滤器属于 Servlet 规范，可以拦截所有符合条件的请求；拦截器属于 Spring MVC，只能拦截进入 DispatcherServlet 的请求。过滤器更底层，拦截器更方便获取控制器和业务信息。')

doc.add_heading('十 高频手撕代码补充', level=1)
p('9 删除倒数第 N 个节点', 'Heading 2')
codeblock('''public ListNode removeNthFromEnd(ListNode head, int n) {
    ListNode dummy = new ListNode(0);
    dummy.next = head;
    ListNode fast = dummy, slow = dummy;
    for (int i = 0; i < n; i++) fast = fast.next;
    while (fast.next != null) { fast = fast.next; slow = slow.next; }
    slow.next = slow.next.next;
    return dummy.next;
}''')
p('10 合并区间', 'Heading 2')
codeblock('''public int[][] merge(int[][] intervals) {
    Arrays.sort(intervals, (a, b) -> Integer.compare(a[0], b[0]));
    List<int[]> list = new ArrayList<>();
    for (int[] cur : intervals) {
        if (list.isEmpty() || list.get(list.size()-1)[1] < cur[0]) list.add(cur);
        else list.get(list.size()-1)[1] = Math.max(list.get(list.size()-1)[1], cur[1]);
    }
    return list.toArray(new int[list.size()][]);
}''')
p('11 移动零', 'Heading 2')
codeblock('''public void moveZeroes(int[] nums) {
    int slow = 0;
    for (int x : nums) if (x != 0) nums[slow++] = x;
    while (slow < nums.length) nums[slow++] = 0;
}''')
p('12 最大子数组和', 'Heading 2')
codeblock('''public int maxSubArray(int[] nums) {
    int best = nums[0], cur = nums[0];
    for (int i = 1; i < nums.length; i++) {
        cur = Math.max(nums[i], cur + nums[i]);
        best = Math.max(best, cur);
    }
    return best;
}''')
p('13 爬楼梯', 'Heading 2')
codeblock('''public int climbStairs(int n) {
    if (n <= 2) return n;
    int a = 1, b = 2;
    for (int i = 3; i <= n; i++) { int c = a + b; a = b; b = c; }
    return b;
}''')
p('14 岛屿数量', 'Heading 2')
codeblock('''public int numIslands(char[][] grid) {
    int count = 0;
    for (int i = 0; i < grid.length; i++)
        for (int j = 0; j < grid[0].length; j++)
            if (grid[i][j] == '1') { count++; dfs(grid, i, j); }
    return count;
}
private void dfs(char[][] g, int i, int j) {
    if (i < 0 || i >= g.length || j < 0 || j >= g[0].length || g[i][j] != '1') return;
    g[i][j] = '0';
    dfs(g,i+1,j); dfs(g,i-1,j); dfs(g,i,j+1); dfs(g,i,j-1);
}''')
p('15 全排列', 'Heading 2')
codeblock('''public List<List<Integer>> permute(int[] nums) {
    List<List<Integer>> ans = new ArrayList<>();
    backtrack(nums, new boolean[nums.length], new ArrayList<>(), ans);
    return ans;
}
private void backtrack(int[] nums, boolean[] used, List<Integer> path, List<List<Integer>> ans) {
    if (path.size() == nums.length) { ans.add(new ArrayList<>(path)); return; }
    for (int i = 0; i < nums.length; i++) if (!used[i]) {
        used[i] = true; path.add(nums[i]); backtrack(nums, used, path, ans);
        path.remove(path.size()-1); used[i] = false;
    }
}''')
p('16 最长递增子序列', 'Heading 2')
codeblock('''public int lengthOfLIS(int[] nums) {
    int[] dp = new int[nums.length];
    Arrays.fill(dp, 1); int ans = 1;
    for (int i = 1; i < nums.length; i++) {
        for (int j = 0; j < i; j++) if (nums[i] > nums[j])
            dp[i] = Math.max(dp[i], dp[j] + 1);
        ans = Math.max(ans, dp[i]);
    }
    return ans;
}''')
p('17 零钱兑换', 'Heading 2')
codeblock('''public int coinChange(int[] coins, int amount) {
    int[] dp = new int[amount + 1];
    Arrays.fill(dp, amount + 1); dp[0] = 0;
    for (int i = 1; i <= amount; i++)
        for (int coin : coins) if (coin <= i)
            dp[i] = Math.min(dp[i], dp[i - coin] + 1);
    return dp[amount] > amount ? -1 : dp[amount];
}''')
p('18 最小栈', 'Heading 2')
codeblock('''class MinStack {
    private final Deque<Integer> data = new ArrayDeque<>();
    private final Deque<Integer> mins = new ArrayDeque<>();
    public void push(int x) { data.push(x); mins.push(mins.isEmpty() ? x : Math.min(x, mins.peek())); }
    public void pop() { data.pop(); mins.pop(); }
    public int top() { return data.peek(); }
    public int getMin() { return mins.peek(); }
}''')
p('19 单例模式 双重检查锁', 'Heading 2')
codeblock('''public class Singleton {
    private static volatile Singleton instance;
    private Singleton() {}
    public static Singleton getInstance() {
        if (instance == null) synchronized (Singleton.class) {
            if (instance == null) instance = new Singleton();
        }
        return instance;
    }
}''')
p('20 生产者消费者', 'Heading 2')
codeblock('''BlockingQueue<Integer> queue = new ArrayBlockingQueue<>(10);
new Thread(() -> { try { queue.put(1); } catch (InterruptedException e) {
    Thread.currentThread().interrupt();
}}).start();
new Thread(() -> { try { System.out.println(queue.take()); } catch (InterruptedException e) {
    Thread.currentThread().interrupt();
}}).start();''')

doc.add_heading('十一 复习与答题技巧', level=1)
bullets(['优先掌握 HashMap、ConcurrentHashMap、JVM 内存与 GC、线程池、MySQL 索引和事务、Spring 事务、Redis 缓存问题。','手撕代码重点刷链表、数组双指针、滑动窗口、二叉树、二分、回溯和动态规划。','写代码前先说思路、复杂度和边界；写完后用空输入、单元素、重复元素和极值进行验证。','项目题要能说明业务背景、技术选型、遇到的问题、解决方案和最终效果。'])
p('推荐复习顺序：Java 基础 → 集合 → 并发 → JVM → MySQL → Spring → Redis → 算法手撕 → 项目场景题。')

doc.core_properties.title='Java 面试八股与手撕代码复习资料'
doc.core_properties.subject='Java 后端面试复习'
doc.core_properties.author=''
doc.save(OUT)
print(OUT)
